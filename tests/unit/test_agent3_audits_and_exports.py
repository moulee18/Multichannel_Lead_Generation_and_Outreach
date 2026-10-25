from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill
import pytest

from backend.app.agents.agent3.qualification_audit import (
    run_independent_qualification_and_geo_audit,
)
from backend.app.agents.agent3.random_audit import run_internal_random_audit
from backend.app.agents.agent3.seed_testing import SeedTestLabService
from backend.app.agents.agent3.submission import evaluate_submission_readiness
from backend.app.database.session import get_session_factory, init_db
from backend.app.exports.csv_exporter import OFFICIAL_LEADS_CSV_COLUMNS, export_leads_csv
from backend.app.exports.workbook_exporter import populate_official_workbook
from backend.app.services.orchestrator import PipelineOrchestrator
from backend.app.utils.errors import IntegrityViolationError


@pytest.fixture()
def db_session(tmp_path: Path):
    db_url = f"sqlite:///{(tmp_path / 'test_agent3.db').as_posix()}"
    init_db(db_url)
    SessionLocal = get_session_factory(db_url)
    with SessionLocal() as session:
        yield session


def test_agent3_audits_geography_and_reproducible_random_sample(db_session):
    orchestrator = PipelineOrchestrator()
    job = orchestrator.run_full_pipeline(
        db=db_session,
        keyword="video podcast",
        source="synthetic_demo",
        run_mode="synthetic_demo",
        max_candidates=35,
    )
    assert job.status == "COMPLETED"

    qual_geo = run_independent_qualification_and_geo_audit(db_session, include_synthetic=True)
    assert qual_geo["total_qualified_checked"] == 30
    assert qual_geo["qualification_pass_rate"] == 100.0
    assert qual_geo["contact_accuracy_rate"] == 100.0
    assert qual_geo["duplicate_count"] == 0

    geo = qual_geo["geographic_audit"]
    assert geo["approved_market_pct"] >= 80.0
    assert geo["represented_countries_count"] >= 3
    assert geo["max_single_country_pct"] <= 60.0
    assert geo["compliant"] is True

    # Reproducible random sample of 20 leads
    sample_a = run_internal_random_audit(db_session, random_seed=42, sample_size=20, include_synthetic=True)
    sample_b = run_internal_random_audit(db_session, random_seed=42, sample_size=20, include_synthetic=True)
    assert sample_a["actual_sample_size"] == 20
    assert [x["creator_id"] for x in sample_a["sampled_leads"]] == [
        x["creator_id"] for x in sample_b["sampled_leads"]
    ]
    assert sample_a["meets_minimum_thresholds"] is True


def test_seed_test_allowlist_tracking_and_replies(db_session):
    # 1. Non-allowlisted real creator address MUST be blocked
    with pytest.raises(IntegrityViolationError):
        SeedTestLabService.schedule_or_execute_seed_test(
            db=db_session,
            channel="email",
            recipient_address="real_youtuber@gmail.com",
            message_preview="Should never send",
            execution_mode="DRY_RUN",
        )

    # 2. Allowlisted owned account succeeds in DRY_RUN and recovery mode
    run = SeedTestLabService.schedule_or_execute_seed_test(
        db=db_session,
        channel="email",
        recipient_address="seedtest.owner@example.org",
        message_preview="Step 1 test",
        execution_mode="DRY_RUN",
        simulate_failure_recovery=True,
    )
    assert run.status == "DRY_RUN_LOGGED"
    assert run.error_recovery_notes is not None

    # 3. Tracking click & reply association
    click_evt = SeedTestLabService.record_tracking_click(db_session, run.tracking_token)
    assert click_evt is not None
    assert click_evt.seed_test_run_id == run.id

    reply_evt = SeedTestLabService.record_reply_event(
        db=db_session,
        seed_test_run_id=run.id,
        sender_address="seedtest.owner@example.org",
        reply_text="Received tracking link!",
    )
    assert reply_evt.seed_test_run_id == run.id


def test_workbook_preservation_and_csv_export(db_session, tmp_path: Path):
    orchestrator = PipelineOrchestrator()
    orchestrator.run_full_pipeline(
        db=db_session,
        keyword="video podcast",
        source="synthetic_demo",
        run_mode="synthetic_demo",
        max_candidates=10,
    )

    csv_path = tmp_path / "leads.csv"
    csv_res = export_leads_csv(db_session, csv_path, include_synthetic=True)
    assert csv_res["record_count"] == 10
    assert csv_path.exists()

    # Build a synthetic Excel workbook template with yellow input cells and green formula cells
    tpl_path = tmp_path / "template.xlsx"
    wb = openpyxl.Workbook()
    ws_leads = wb.active
    ws_leads.title = "Leads"
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    for idx, col_name in enumerate(OFFICIAL_LEADS_CSV_COLUMNS, start=1):
        ws_leads.cell(row=1, column=idx, value=col_name)

    # Add a green formula cell in column 25 that must be preserved
    ws_leads.cell(row=1, column=25, value="formula_check")
    f_cell = ws_leads.cell(row=2, column=25, value="=COUNTA(A2:A100)")
    f_cell.fill = green_fill

    for r in range(2, 12):
        c_cell = ws_leads.cell(row=r, column=1)
        c_cell.fill = yellow_fill

    ws_sc = wb.create_sheet("Self-check")
    ws_sc.cell(row=1, column=1, value="Claim")
    ws_sc.cell(row=1, column=2, value="Evidence Link")
    wb.save(tpl_path)

    out_wb_path = tmp_path / "completed_workbook.xlsx"
    wb_res = populate_official_workbook(db_session, tpl_path, out_wb_path, csv_path, include_synthetic=True)
    assert wb_res["status"] == "COMPLETED"
    assert wb_res["consistent_with_leads_csv"] is True

    wb_check = openpyxl.load_workbook(out_wb_path)
    assert wb_check["Leads"].cell(row=2, column=25).value == "=COUNTA(A2:A100)"
    assert wb_check["Leads"].cell(row=2, column=1).value is not None
    assert wb_check["Self-check"].cell(row=2, column=2).value.startswith("http")

    # Readiness check must report NOT_READY because synthetic demo leads are never counted toward 1,000 genuine leads
    readiness = evaluate_submission_readiness(db_session)
    assert readiness["overall_status"] == "NOT_READY"
    assert readiness["genuine_qualified_leads"] == 0
    assert readiness["synthetic_demo_qualified_leads"] == 10
